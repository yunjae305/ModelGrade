"""새 질문 리스트 만들기 — 질문생성 프롬프트(v3.6)를 Claude Code CLI 에 맡겨 질문 · 첨부 · 긴 문서를 세트 폴더에 바로 만들고,
우리 검사기로 통과할 때까지 이어서 고친다 (사용자 10-26 저녁: Codex 대신 Claude Code — Windows 에서 Codex 는 폴더에 쓰지 못함).

    python -X utf8 tools/new_question_set.py make [--size full|small]              # modelgrade start 의 2번과 같다
    python -X utf8 tools/new_question_set.py check <세트 폴더> [--size …] [--json]  # 검사만 — Claude 도 작업 중에 부른다

흐름: 세트 폴더(참고자료/질문세트_<날짜_시각>)를 만들고 프롬프트를 복사(세트 이름 · 크기만 바꿈) → Claude 가 그 폴더 안에서 작업
(파일 쓰기는 그 폴더만, 명령은 python 만) → 회차가 끝날 때마다 검사 → 통과하면 set.json · criteria_catalog.json 을 맞추고 끝,
아니면 문제 목록을 주고 다음 회차. 멈춰도 같은 명령이면 만들던 폴더(생성중.json 이 있는 가장 최근 폴더)에서 이어서 한다.
검사가 통과하기 전에는 그 세트로 수집하지 않는다(k 범위 확인 뒤에만 수집).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import threading
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

PROMPT_PATH = ROOT / "참고자료" / "프롬프트" / "질문생성_프롬프트.md"
SETS_DIR = ROOT / "참고자료"
MARKER = "생성중.json"          # 만드는 중인 세트 — 끝나면 생성기록/생성.json 으로 옮긴다
LOG_DIR = "생성기록"
GENERATOR = "question-gen-v3.6 · Claude Code"
DEFAULT_MODEL = "claude-opus-5-5"
DEFAULT_EFFORT = "high"
MAX_ROUNDS = 8
IDLE_TIMEOUT = 40 * 60          # 이만큼 아무 출력이 없으면 그 회차를 멈추고 다음 회차로 (배경 작업은 시간 제한, 10-26)
ROUND_TIMEOUT = 8 * 3600
#: Claude 가 쓸 수 있는 도구 — 파일 쓰기는 acceptEdits(작업 폴더 안만), 명령은 python 만. Windows 의 Claude Code 는 PowerShell 도구로 실행한다
ALLOWED_TOOLS = ("Read", "Glob", "Grep", "TodoWrite", "Agent", "Task",
                 "Bash(python:*)", "Bash(python3:*)", "PowerShell(python:*)", "PowerShell(python3:*)")

#: 세트 크기 — full 은 프롬프트 [요청] 그대로(78), small 은 흐름 시험용 5개
SPECS: Dict[str, Dict[str, Any]] = {
    "full": {"total": 78, "multi": 11, "long": 10, "empty": 10, "attachment": 12,
             "attachment_types": {"txt": 3, "docx": 3, "pdf": 3, "image": 3}, "min_applicable": 10},
    "small": {"total": 5, "multi": 0, "long": 0, "empty": 1, "attachment": 2,
              "attachment_types": {"docx": 1, "pdf": 1}, "min_applicable": 0},
}
SMALL_REQUEST = """# [요청] — 사람이 바꾸는 부분
- 만들 질문 수: **5개**(흐름 시험용 작은 세트) = 단일턴 5개 — 일반 2개 · 첨부 2개(docx 1개 · pdf 1개) · 빈 출력 1개. 10턴 대화 · 긴 문서 질문 · reference.md 는 만들지 않습니다.
- 기준별 최소 10개 · 대화 · 긴 문서에 관한 규칙은 이번에는 적용하지 않습니다. 나머지 규칙(k 6~10, 함정, 첨부 content, questions.json 형식)은 그대로입니다.
- 세트 이름: `{set_id}`.

"""
MULTI_TURN_TOKENS = 10_000      # 턴마다 질문만으로 (프롬프트 5장)
MULTI_TOTAL_TOKENS = 100_000
CONV_K = (6.0, 10.6)            # 대화 평균 k — 반올림 전, 10.6 까지 허용(사용자 10-21 밤)


# --- 검사 -----------------------------------------------------------------------------------------------------------------

def _norm(text: str) -> str:
    return re.sub(r"\s+", "", str(text or ""))


def _ids(xs: List[Any], limit: int = 12) -> str:
    xs = list(xs)
    return ", ".join(f"Q{x}" for x in xs[:limit]) + (f" 외 {len(xs) - limit}개" if len(xs) > limit else "")


def _file_lines(path: Path, kind: str) -> Optional[List[str]]:
    """파일에서 읽은 글 줄(대조용) — 읽을 수 없으면 None(대조 생략)."""
    if kind == "docx":
        from docx import Document
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines += [c.text for c in row.cells]
        return lines
    if kind == "pdf":   # pypdfium2 는 필수 의존성(판정의 PDF 쪽 이미지) — 깨진 PDF 는 여기서 예외 → '파일을 열지 못함'
        import pypdfium2
        doc = pypdfium2.PdfDocument(str(path))
        try:
            return "\n".join(doc[i].get_textpage().get_text_range() for i in range(len(doc))).splitlines()
        finally:
            doc.close()
    return None


def check_attachment(set_dir: Path, a: Any) -> List[str]:
    """첨부 하나 — 이름 · 종류 · 파일이 열리는가 · content 가 파일 글과 같은가."""
    from webapp.attachment_files import ATTACHMENT_TYPES
    if not isinstance(a, dict):
        return ["첨부 항목이 객체가 아님"]
    name, kind, path, content = a.get("name"), a.get("type"), str(a.get("path") or ""), a.get("content")
    out: List[str] = []
    if kind not in ATTACHMENT_TYPES:
        out.append(f"{name}: type 이 {kind!r} (txt · docx · pdf · image 중 하나)")
    if not path or Path(path).name != path:
        return out + [f"{name}: path 는 files/ 를 뺀 파일 이름 하나({path!r})"]
    f = set_dir / "files" / path
    if not f.is_file() or f.stat().st_size == 0:
        return out + [f"{path}: files/ 에 파일이 없음"]
    if not isinstance(content, str) or not content.strip():
        return out + [f"{path}: content(파일에 실제로 들어간 글 전체)가 비어 있음"]
    if "\\n" in content:
        out.append(f"{path}: content 에 역슬래시-n 두 글자가 그대로 있음(줄바꿈은 JSON 의 \\n 이스케이프로)")
    try:
        if kind == "txt":
            if _norm(f.read_text(encoding="utf-8")) != _norm(content):
                out.append(f"{path}: 파일 글과 content 가 다름")
        elif kind == "image":
            from PIL import Image
            with Image.open(f) as im:
                im.verify()
        elif kind in ("docx", "pdf"):
            if kind == "pdf" and not f.read_bytes()[:5].startswith(b"%PDF"):
                return out + [f"{path}: PDF 파일이 아님"]
            flat = lambda t: _norm(t).replace("|", "")   # 표 칸 막대는 빼고 — PDF 글 뽑기는 칸 글을 구분자 없이 이어 붙인다
            lines = [flat(x) for x in (_file_lines(f, kind) or []) if len(flat(x)) >= 4]
            body = flat(content)
            if not lines:   # 빈 문서 · 글을 뽑을 수 없는 문서 — content 와 맞는지 확인할 수 없다 (Codex 검토 10-26 밤)
                out.append(f"{path}: 파일에서 글을 읽지 못함(빈 문서이거나 글이 그림으로만 들어감) — content 의 글이 실제로 든 파일로 다시 만들기")
            else:
                missing = [x for x in lines if x not in body]
                if len(missing) > 0.1 * len(lines):
                    out.append(f"{path}: 파일 글 {len(lines)}줄 중 {len(missing)}줄이 content 에 없음(예: {missing[0][:30]}) — 파일을 다시 열어 content 와 맞추기")
    except Exception as exc:   # 파일이 깨졌다 — 다시 만들게
        out.append(f"{path}: 파일을 열지 못함 — {type(exc).__name__}: {str(exc)[:80]}")
    return out


def check_set(set_dir: Path, size: str = "full") -> Dict[str, Any]:
    """세트 폴더 검사 → {"ok", "problems": [고칠 것 한 줄씩], "counts", "k", "conversations", "short"}. 평가 때와 같은 코드(질문 인식기 ·
    k_var · 적용 수 · 토큰 수)로 센다 — 모델은 부르지 않는다. 만든 질문의 모양이 이상해 검사가 넘어지면 그 오류를 '고칠 것'으로 돌려준다
    (회차 반복이 멈추지 않고 Claude 가 다음 회차에 고치게 — Codex 검토 10-26 밤)."""
    try:
        return _check_set(set_dir, size)
    except Exception as exc:
        return {"ok": False, "size": size, "counts": {},
                "problems": [f"검사가 질문 모양 때문에 멈춤 — {type(exc).__name__}: {str(exc)[:200]} (questions.json 의 키 · 값 모양을 프롬프트 9장대로)"]}


def _check_set(set_dir: Path, size: str) -> Dict[str, Any]:
    spec = SPECS[size]
    res: Dict[str, Any] = {"ok": False, "size": size, "problems": [], "counts": {}}
    problems: List[str] = res["problems"]
    qpath = set_dir / "questions.json"
    if not qpath.is_file():
        problems.append("questions.json 이 아직 없음")
        return res
    try:
        groups = json.loads(qpath.read_text(encoding="utf-8"))
    except ValueError as exc:
        problems.append(f"questions.json 을 JSON 으로 읽지 못함 — {exc}")
        return res
    if not isinstance(groups, list) or not groups or not all(isinstance(g, dict) for g in groups):
        problems.append("questions.json 은 질문 객체의 목록이어야 함")
        return res
    bad = []
    for i, g in enumerate(groups, 1):
        turns = g.get("turns")
        if (g.get("id") != i or not isinstance(turns, list) or not turns or not isinstance(g.get("target_items"), list)
                or [t.get("turn") if isinstance(t, dict) else None for t in turns] != list(range(1, len(turns) + 1))
                or not all(isinstance(t.get("question"), str) and t["question"].strip() for t in turns)
                or not all(isinstance(t.get("attachments") or [], list) and isinstance(t.get("reference") or "", str) for t in turns)):
            bad.append(g.get("id", f"{i}번째"))
    if bad:
        problems.append(f"형식이 틀린 질문 {len(bad)}개({', '.join(map(str, bad[:12]))}) — id 는 1부터 차례로, turns 는 turn 1부터 · question 글, "
                        "target_items 목록 (프롬프트 9장)")
        return res

    from tools.check_question_set import check_items
    from webapp.question_generation import flatten, k_var_by_group, k_var_summary, k_var_units, target_counts
    report = check_items(flatten(groups), pipeline="items", minimum=max(1, spec["min_applicable"]))
    ksum = k_var_summary(groups, report)
    kmap = k_var_by_group(groups, report)
    singles = [g for g in groups if len(g["turns"]) == 1]
    multi = [g for g in groups if len(g["turns"]) == 10]
    odd = [g["id"] for g in groups if len(g["turns"]) not in (1, 10)]
    long_ids = [g["id"] for g in singles if isinstance(g["turns"][0].get("reference"), str) and g["turns"][0]["reference"].strip()]
    att = [g for g in singles if g["turns"][0].get("attachments")]
    empty_ids = list(ksum.get("exempt_ids") or [])
    kinds: Dict[str, int] = {}
    for g in att:
        for a in g["turns"][0]["attachments"]:
            if isinstance(a, dict):
                kinds[str(a.get("type"))] = kinds.get(str(a.get("type")), 0) + 1
    counts = {"total": len(groups), "single": len(singles), "multi": len(multi), "long": len(long_ids), "attachment": len(att),
              "empty": len(empty_ids), "attachment_types": kinds, "turns": sum(len(g["turns"]) for g in groups)}
    res["counts"] = counts

    names = {"total": "질문", "multi": "10턴 대화", "long": "긴 문서 질문", "empty": "빈 출력 질문", "attachment": "첨부 질문"}
    for key, label in names.items():
        if counts[key] != spec[key]:
            problems.append(f"구성: {label} {counts[key]}개 (목표 {spec[key]}개)")
    if kinds != spec["attachment_types"]:
        problems.append(f"구성: 첨부 종류 {kinds or '없음'} (목표 {spec['attachment_types']})")
    if odd:
        problems.append(f"구성: 1턴도 10턴도 아닌 질문 {_ids(odd)} — 대화는 정확히 10턴, 2턴 대화 금지")
    for g in groups:
        for t in g["turns"]:
            for a in t.get("attachments") or []:
                where = f"Q{g['id']}" + (f" {t['turn']}턴" if len(g["turns"]) > 1 else "")
                problems += [f"첨부 {where} {p}" for p in check_attachment(set_dir, a)]

    out_k = [f"Q{i}(k={kmap[i]['k']})" for i in ksum["out_of_range_ids"]]
    if out_k:
        problems.append(f"k 범위(6~10) 밖 단일 질문 {len(out_k)}개: {', '.join(out_k[:15])}" + (" …" if len(out_k) > 15 else ""))
    res["k"] = ksum

    units = k_var_units()
    by_case = {r["case_id"]: r for r in report.get("cases") or []}
    from code_eval.long_context import MIN_TOKENS, count_tokens
    convs: Dict[str, Any] = {}
    for g in multi:
        ks, tks = [], []
        for t in g["turns"]:
            r = by_case.get(f"Q{g['id']}_T{t['turn']}") or {}
            ks.append(len(set((r.get("checks") or []) + (r.get("items") or [])) & units))
            tks.append(count_tokens(t["question"]))
        avg = round(sum(ks) / len(ks), 2)
        convs[str(g["id"])] = {"turn_k": ks, "avg_k": avg, "turn_tokens": tks, "total_tokens": sum(tks)}
        if not CONV_K[0] <= avg <= CONV_K[1]:
            problems.append(f"대화 Q{g['id']}: 턴 평균 k {avg} (6~10.6) — 턴별 {ks}")
        short_turns = [i + 1 for i, n in enumerate(tks) if n < MULTI_TURN_TOKENS]
        if short_turns or sum(tks) < MULTI_TOTAL_TOKENS:
            problems.append(f"대화 Q{g['id']}: 질문만으로 {MULTI_TURN_TOKENS:,}토큰 미만인 턴 {short_turns or '없음'} · "
                            f"전체 {sum(tks):,}토큰(목표 {MULTI_TOTAL_TOKENS:,} 이상)")
    res["conversations"] = convs

    cache: Dict[str, int] = {}
    for i in long_ids:
        t = groups[i - 1]["turns"][0]
        h = hashlib.sha1(t["reference"].encode("utf-8")).hexdigest()
        if h not in cache:
            cache[h] = count_tokens(t["reference"])
        n = cache[h] + count_tokens(t["question"])
        if n < MIN_TOKENS:
            problems.append(f"긴 문서 Q{i}: 질문 + reference {n:,}토큰 (목표 {MIN_TOKENS:,} 이상)")

    if spec["min_applicable"]:
        tc = target_counts(report)
        short = {k: v for k, v in sorted(tc.items(), key=lambda x: (x[1], x[0])) if v < spec["min_applicable"]}
        res["short"] = short
        if short:
            problems.append(f"적용 질문이 {spec['min_applicable']}개 미만인 기준 {len(short)}개: "
                            + ", ".join(f"{k} {v}개" for k, v in list(short.items())[:20]) + (" …" if len(short) > 20 else ""))
    for key, label in (("errors", "인식기 오류"), ("duplicates", "같은 질문")):
        xs = report.get(key) or []
        if xs:
            problems.append(f"{label} {len(xs)}건: {json.dumps(xs[:3], ensure_ascii=False)[:200]}")
    res["ok"] = not problems
    return res


def render(res: Dict[str, Any]) -> str:
    c = res.get("counts") or {}
    head = (f"질문 {c.get('total', 0)}개 (단일 {c.get('single', 0)} · 10턴 대화 {c.get('multi', 0)} · 긴 문서 {c.get('long', 0)} · "
            f"첨부 {c.get('attachment', 0)} · 빈 출력 {c.get('empty', 0)})") if c else ""
    if res.get("ok"):
        k = res.get("k") or {}
        return f"검사: 통과 — {head} · 단일 질문 k 분포 {k.get('by_k')}"
    return "\n".join([f"검사: 고칠 것 {len(res['problems'])}가지" + (f" — {head}" if head else "")] + [f"- {p}" for p in res["problems"]])


# --- 만들기 ---------------------------------------------------------------------------------------------------------------

def request_block(prompt: str, set_id: str, size: str) -> str:
    """프롬프트의 [요청] 블록 — full 은 세트 이름만, small 은 블록을 시험용으로 바꾼다."""
    start = prompt.index("# [요청]")
    end = prompt.index("\n# 0.", start) + 1
    if size == "small":
        block = SMALL_REQUEST.format(set_id=set_id)
    else:
        block = re.sub(r"세트 이름: `[^`]+`", f"세트 이름: `{set_id}`", prompt[start:end])
    return prompt[:start] + block + prompt[end:]


def task_text(set_dir: Path, size: str, round_no: int, last_check: Optional[str]) -> str:
    from webapp.attachment_files import font_path
    try:
        font = font_path()
    except Exception:
        font = "시스템 한글 글꼴"
    checker = (ROOT / "tools" / "new_question_set.py").as_posix()
    lines = [
        f"이 폴더({set_dir.name})에 한국어 AI 답변 평가용 질문 세트를 만듭니다. 할 일은 이 폴더의 `질문생성_프롬프트.md` 에 모두 있습니다 — "
        "처음부터 끝까지 읽고 그대로 따르세요.",
        "이 실행 환경에 맞춰 아래만 다르게 합니다.",
        "- 결과는 ZIP 대신 이 폴더에 바로 씁니다: questions.json · files/ · set.json · README.md (긴 문서가 있으면 reference.md). ZIP 은 만들지 않습니다.",
        "- 파일은 이 폴더 안에만 씁니다. 명령은 `python -X utf8 스크립트.py …` 처럼 python 한 줄로만 실행합니다"
        "(cd · && · ; · 파이프 · 환경 변수 앞붙임 없이 — 작업 폴더가 이미 이 폴더입니다).",
        "- 긴 글(턴 문서 · reference.md · 첨부 본문)은 여러 파일로 나눠 쓰고 python 스크립트로 이어 붙입니다. 같은 문단 반복 · 채움 글은 쓰지 않습니다. "
        "오래 걸리는 부분은 보조 에이전트(Agent)에 나눠 맡겨도 됩니다.",
        f"- 첨부: DOCX 는 python-docx, PDF 는 fpdf2, 이미지는 Pillow 로 만듭니다(설치돼 있음, 한글 글꼴 {font}). 문서 만들기 스킬이 있으면 써도 됩니다.",
        "- 단계가 끝날 때마다 `진행상황.md` 에 한 일 · 남은 일을 적습니다. 멈추면 다음 회차가 이 파일을 보고 이어서 합니다.",
        f"- 공식 검사: `python -X utf8 {checker} check . --size {size}` — 구성 · 첨부 · 길이 · k · 기준별 적용 수를 실제 평가 코드로 셉니다"
        "(프롬프트의 check_question_set.py 점검도 이 안에 들어 있음). 질문을 고칠 때마다 다시 돌리고, 이 검사가 '통과'일 때만 끝냅니다.",
        "- 사용자에게 묻지 말고 끝까지 진행합니다. 마지막 답은 검사 결과 한두 줄만.",
    ]
    if round_no > 1:
        lines += ["", f"이번은 {round_no}회차입니다. 앞 회차 작업이 이 폴더에 있습니다 — `진행상황.md` 와 아래 검사 결과를 먼저 보고, "
                      "처음부터 다시 만들지 말고 이어서 고치세요.",
                  "", "[지난 검사 결과]", last_check or "(검사 결과 없음 — questions.json 부터 이어서)"]
    return "\n".join(lines)


def _read_json(path: Path) -> Dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def _write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")


def show(path: Path) -> str:
    return path.relative_to(ROOT).as_posix() if path.is_relative_to(ROOT) else str(path)


def unfinished_set() -> Optional[Path]:
    """만들다 멈춘 세트 — 생성중.json 이 있는 가장 최근 폴더."""
    found = sorted((p.parent for p in SETS_DIR.glob(f"질문세트_*/{MARKER}")), key=lambda p: p.stat().st_mtime, reverse=True)
    return found[0] if found else None


def made_sets() -> List[Dict[str, Any]]:
    """이 프로그램이 만든 완성 세트 — 최근 것부터 [{dir, set_id, created_at, total}]."""
    out = []
    for p in SETS_DIR.glob("질문세트_*/set.json"):
        s = _read_json(p)
        if str(s.get("generator") or "").startswith(GENERATOR) and s.get("status") == "COMPLETED" and not (p.parent / MARKER).exists():
            out.append({"dir": p.parent, "set_id": s.get("set_id"), "created_at": str(s.get("created_at") or ""),
                        "total": (s.get("counts") or {}).get("total")})
    return sorted(out, key=lambda x: x["created_at"], reverse=True)


def new_set_dir(now: Optional[datetime] = None) -> Path:
    now = now or datetime.now()
    d = SETS_DIR / f"질문세트_{now:%Y-%m-%d_%H%M}"
    n = 2
    while d.exists():
        d = SETS_DIR / f"질문세트_{now:%Y-%m-%d_%H%M}_{n}"
        n += 1
    return d


def criteria_catalog_file() -> Dict[str, Any]:
    """세트 폴더의 criteria_catalog.json — 결과 보고서가 '언제 보나'(항상 · 답에 따라 · 질문이 요구할 때)를 여기서 읽는다."""
    from code_eval import criteria_registry as reg
    from code_eval.semantic_items import ITEMS
    from webapp.question_generation import criteria_catalog, k_var_units
    cat = criteria_catalog()
    code_all = {u[0] for u in reg.CODE_UNITS if u[2] == "all"}
    return {"source": "ModelGrade", "registry": reg.REGISTRY_VERSION, "catalog": cat,
            "always_ids": sorted(code_all | {i.id for i in ITEMS if i.always}),
            "answer_dependent_ids": sorted({u[0] for u in reg.CODE_UNITS if u[2] == "answer"} | {"J5-3.a"}),
            "variable_ids": sorted(k_var_units(cat))}


def claude_version(exe: str) -> str:
    try:
        return subprocess.run([exe, "--version"], capture_output=True, text=True, timeout=30).stdout.strip().split()[0]
    except Exception:
        return "?"


def finalize(set_dir: Path, meta: Dict[str, Any], res: Dict[str, Any]) -> None:
    """통과한 세트 — set.json 을 우리 형식으로 맞추고 criteria_catalog.json 을 쓰고, 생성중 표시를 생성기록으로 옮긴다."""
    s = _read_json(set_dir / "set.json")
    c = res["counts"]
    s.update({"id": meta["set_id"], "set_id": meta["set_id"], "created_at": s.get("created_at") or meta["started_at"],
              "counts": {"total": c["total"], "single": c["single"], "multi_turn": c["multi"], "long_context": c["long"],
                         "attachment": c["attachment"], "empty_output": c["empty"]},
              "generator": f"{GENERATOR} {meta.get('claude_version', '?')} · {meta['model']} {meta['effort']}",
              "status": "COMPLETED", "size": meta["size"], "k_var": res.get("k"),
              "multiturn_k_exact": {q: v["avg_k"] for q, v in (res.get("conversations") or {}).items()},
              "precheck_min_applicable": SPECS[meta["size"]]["min_applicable"], "rounds": len(meta.get("rounds") or [])})
    _write_json(set_dir / "set.json", s)
    _write_json(set_dir / "criteria_catalog.json", criteria_catalog_file())
    meta["finished_at"] = datetime.now().isoformat(timespec="seconds")
    _write_json(set_dir / LOG_DIR / "생성.json", meta)
    (set_dir / MARKER).unlink(missing_ok=True)


def describe_tool(block: Dict[str, Any]) -> str:
    """Claude 의 도구 호출 하나 → 진행 판 메모 한 줄."""
    name, inp = block.get("name"), block.get("input") or {}
    if name in ("Write", "Edit", "MultiEdit"):
        return f"{'쓰는 중' if name == 'Write' else '고치는 중'} {Path(str(inp.get('file_path') or '')).name}"
    if name in ("Bash", "PowerShell"):
        cmd = " ".join(str(inp.get("command") or "").split())
        if "new_question_set.py" in cmd and " check" in cmd:
            return "검사 돌리는 중"
        m = re.search(r"([\w가-힣.-]+\.py)\b", cmd)
        return f"실행 {m.group(1)}" if m else "python 코드 실행" if re.search(r"\s-c\s", f" {cmd} ") else f"실행 {cmd[:40]}"
    if name in ("Agent", "Task"):
        return f"보조 작업 — {str(inp.get('description') or '')[:40]}"
    if name in ("Read", "Glob", "Grep"):
        return "읽는 중"
    return str(name)


def run_round(exe: str, set_dir: Path, task: str, meta: Dict[str, Any], round_no: int,
              on_note: Callable[[str], None]) -> Dict[str, Any]:
    """Claude 한 회차 — stream-json 을 생성기록/claude-<회차>.jsonl 에 남기고, 도구 호출을 진행 판 메모로."""
    logs = set_dir / LOG_DIR
    logs.mkdir(parents=True, exist_ok=True)
    cmd = [exe, "-p", "--output-format", "stream-json", "--verbose", "--model", meta["model"], "--effort", meta["effort"],
           "--permission-mode", "acceptEdits", "--allowedTools", *ALLOWED_TOOLS]
    env = dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8")
    out: Dict[str, Any] = {"round": round_no, "started": datetime.now().isoformat(timespec="seconds"), "result": None, "error": None}
    last = [time.monotonic()]
    with (logs / f"claude-{round_no}.jsonl").open("a", encoding="utf-8") as raw, \
            (logs / f"claude-{round_no}.err").open("a", encoding="utf-8") as err:
        proc = subprocess.Popen(cmd, cwd=str(set_dir), stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=err,
                                text=True, encoding="utf-8", errors="replace", env=env, start_new_session=os.name != "nt")
        try:
            proc.stdin.write(task)
            proc.stdin.close()
        except OSError:   # 바로 죽었다 — 아래에서 종료 코드로 남는다
            pass

        def pump() -> None:
            for line in proc.stdout:
                raw.write(line)
                raw.flush()
                last[0] = time.monotonic()
                try:
                    ev = json.loads(line)
                except ValueError:
                    continue
                if ev.get("type") == "assistant":
                    for block in (ev.get("message") or {}).get("content") or []:
                        if isinstance(block, dict) and block.get("type") == "tool_use":
                            on_note(describe_tool(block))
                elif ev.get("type") == "result":
                    out["result"] = {"subtype": ev.get("subtype"), "is_error": ev.get("is_error"), "text": str(ev.get("result") or "")[:2000],
                                     "session_id": ev.get("session_id"), "num_turns": ev.get("num_turns")}

        t = threading.Thread(target=pump, daemon=True)
        t.start()
        t0 = time.monotonic()
        try:
            while proc.poll() is None:
                time.sleep(2)
                idle = time.monotonic() - last[0] > IDLE_TIMEOUT
                if idle or time.monotonic() - t0 > ROUND_TIMEOUT:
                    out["error"] = (f"출력 없음 {IDLE_TIMEOUT // 60}분" if idle else "회차 시간 초과") + " — 멈추고 다음 회차"
                    kill_tree(proc)
                    break
        except KeyboardInterrupt:   # Ctrl+C — Claude 를 뒤에 남기지 않는다(만들던 폴더는 그대로라 다시 2번이면 이어서)
            kill_tree(proc)
            raise
        proc.wait()
        t.join(timeout=10)
    out["exit"] = proc.returncode
    out["ended"] = datetime.now().isoformat(timespec="seconds")
    return out


def kill_tree(proc: subprocess.Popen) -> None:
    """Claude 와 그 아래에서 돌던 python 스크립트 · 보조 에이전트까지 멈춘다."""
    try:
        if os.name == "nt":
            subprocess.run(["taskkill", "/T", "/F", "/PID", str(proc.pid)], capture_output=True, timeout=30)
        else:
            import signal
            os.killpg(proc.pid, signal.SIGKILL)
    except Exception:
        pass
    try:
        proc.kill()
    except OSError:
        pass


LIMIT_HINT = re.compile(r"usage limit|rate limit|limit (?:will )?reset|한도", re.I)


def make(size: str = "full", model: Optional[str] = None, effort: Optional[str] = None, max_rounds: int = MAX_ROUNDS,
         resume: bool = True, set_dir: Optional[Path] = None, say: Callable[[str], None] = lambda s: print(s, flush=True)) -> Path:
    """새 질문 리스트 → 검사를 통과한 세트 폴더. 만들던 폴더가 있으면(resume) 거기서 이어서."""
    from code_eval.claude_cli_client import find_claude_executable
    from tools.live_board import LiveBoard
    from tools import setup_wizard
    exe = find_claude_executable()
    if not exe:
        raise SystemExit("Claude Code CLI(claude)를 찾지 못했습니다 — modelgrade setup 으로 설치 · 로그인하세요")
    if not setup_wizard.claude_logged_in(exe):   # 회차를 돌리기 전에 — 로그인 안 된 채 시작하면 두 회차를 헛돈다
        if setup_wizard.interactive():
            say("Claude Code 로그인이 필요합니다 — 브라우저가 열립니다")
            subprocess.run([exe, "auth", "login"])
        if not setup_wizard.claude_logged_in(exe):
            raise SystemExit("Claude Code 에 로그인하지 않아 질문 리스트를 만들 수 없습니다 — claude auth login 뒤 다시 실행하세요")
    set_dir = set_dir or (unfinished_set() if resume else None)
    if set_dir and (set_dir / MARKER).exists():
        meta = _read_json(set_dir / MARKER)
        say(f"만들던 질문 리스트를 이어서 만듭니다 — {show(set_dir)}")
    else:
        set_dir = set_dir or new_set_dir()
        set_dir.mkdir(parents=True, exist_ok=True)
        now = datetime.now()
        set_id = f"qs-{now:%Y-%m-%d-%H%M}" + ("-small" if size == "small" else "")
        meta = {"set_id": set_id, "size": size, "model": model or DEFAULT_MODEL, "effort": effort or DEFAULT_EFFORT,
                "started_at": now.isoformat(timespec="seconds"), "prompt": show(PROMPT_PATH), "rounds": []}
        (set_dir / "질문생성_프롬프트.md").write_text(request_block(PROMPT_PATH.read_text(encoding="utf-8"), set_id, size), encoding="utf-8")
        _write_json(set_dir / MARKER, meta)
        say(f"질문 리스트를 만듭니다 — {show(set_dir)} (질문 {SPECS[size]['total']}개 · Claude Code {meta['model']} · "
            "오래 걸립니다, 멈춰도 같은 메뉴로 이어서)")
    meta["claude_version"] = claude_version(exe)
    spec = SPECS[meta["size"]]

    res = check_set(set_dir, meta["size"]) if (set_dir / "questions.json").exists() else None
    if res and res["ok"]:
        finalize(set_dir, meta, res)
        return set_dir
    board = LiveBoard({"gen": "질문 리스트"}, {"gen": spec["total"]})
    t_start = time.monotonic()
    state: Dict[str, Any] = {"note": "시작", "round": len(meta["rounds"]) + 1, "stop": False}

    def done_count() -> int:
        try:
            g = json.loads((set_dir / "questions.json").read_text(encoding="utf-8"))
            return len(g) if isinstance(g, list) else 0
        except (OSError, ValueError):
            return 0

    def redraw(note: Optional[str] = None) -> None:
        if note:
            state["note"] = note
        mins = int((time.monotonic() - t_start) // 60)
        board.update("gen", done=done_count(), note=f"{state['round']}회차 · {mins}분 · {state['note']}")

    def ticker() -> None:
        while not state["stop"]:
            time.sleep(30)
            if not state["stop"]:
                redraw()

    board.start()
    threading.Thread(target=ticker, daemon=True).start()
    try:
        for _ in range(max_rounds):
            n = len(meta["rounds"]) + 1
            state["round"] = n
            redraw("Claude 작업 중")
            r = run_round(exe, set_dir, task_text(set_dir, meta["size"], n, render(res) if res else None), meta, n, redraw)
            redraw("검사 중")
            res = check_set(set_dir, meta["size"])
            r["check_ok"], r["problems"] = res["ok"], res["problems"][:40]
            meta["rounds"].append(r)
            _write_json(set_dir / MARKER, meta)
            _write_json(set_dir / LOG_DIR / f"검사-{n}.json", res)
            if res["ok"]:
                redraw("검사 통과")
                break
            result = r.get("result") or {}
            if result.get("is_error") and LIMIT_HINT.search(result.get("text") or ""):
                board.log(f"  ⚠ Claude 사용 한도 — {(result.get('text') or '')[:120]}")
                break
            if r.get("exit") and not result.get("num_turns"):
                fails = [x for x in meta["rounds"][-2:] if x.get("exit") and not (x.get("result") or {}).get("num_turns")]
                if len(fails) >= 2:   # 두 번 연속 시작도 못 함 — 같은 실패를 되풀이하지 않는다
                    board.log(f"  ⚠ Claude Code 가 실행되지 않습니다 — {show(set_dir / LOG_DIR / f'claude-{n}.err')}")
                    break
            board.log(f"  {n}회차 끝 — 고칠 것 {len(res['problems'])}가지, 이어서 고칩니다")
    finally:
        state["stop"] = True
        board.close()
    if not res or not res["ok"]:
        raise SystemExit("질문 리스트가 아직 검사를 통과하지 못했습니다 — 같은 메뉴(2)로 다시 실행하면 이어서 만듭니다.\n" + (render(res) if res else ""))
    finalize(set_dir, meta, res)
    return set_dir


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="새 질문 리스트 만들기 · 검사 (질문생성 프롬프트 + Claude Code CLI)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    mk = sub.add_parser("make", help="새 질문 리스트 만들기 — 만들던 것이 있으면 이어서")
    mk.add_argument("--size", choices=list(SPECS), default="full", help="full = 78개(프롬프트 그대로) · small = 흐름 시험용 5개")
    mk.add_argument("--model", help=f"Claude 모델 (기본 {DEFAULT_MODEL})")
    mk.add_argument("--effort", help=f"추론 노력 (기본 {DEFAULT_EFFORT})")
    mk.add_argument("--rounds", type=int, default=MAX_ROUNDS, help="검사를 통과할 때까지 이어서 고치는 최대 회차")
    mk.add_argument("--new", action="store_true", help="만들던 것이 있어도 새로 시작")
    ck = sub.add_parser("check", help="세트 폴더 검사 (구성 · 첨부 · 길이 · k · 기준별 적용 수)")
    ck.add_argument("dir", type=Path)
    ck.add_argument("--size", choices=list(SPECS), default=None, help="기본: 폴더의 생성중.json · set.json 의 size, 없으면 full")
    ck.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    if a.cmd == "check":
        d = a.dir.resolve()
        size = a.size or _read_json(d / MARKER).get("size") or _read_json(d / "set.json").get("size") or "full"
        res = check_set(d, size if size in SPECS else "full")
        print(json.dumps(res, ensure_ascii=False, indent=1, default=list) if a.json else render(res))
        return 0 if res["ok"] else 1
    d = make(a.size, a.model, a.effort, a.rounds, resume=not a.new)
    print(f"질문리스트가 완성되었습니다 — {show(d)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
