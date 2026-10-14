"""질문 생성용 첨부파일 만들기 — GPT 가 쓴 첨부 명세로 TXT·DOCX·PDF·이미지(PNG) 파일을 만든다 (J8, 2026-10-13 사용자 결정).

명세 {name, type, content, chart?}:
- ``content`` 는 파일에 들어갈 내용 전체다. 줄 단위 표기: '# 제목'·'## 소제목', '- 항목', '| a | b |' 표 행, '[다음 쪽]' 쪽 나눔(PDF·DOCX).
- 이미지는 ``chart``={title, labels, values, unit} 이면 막대그래프, 아니면 content 의 글·표를 그린다.
만든 파일의 실제 내용(평가 모델이 받는 ``content``)은 그린 그대로의 글로 남긴다 — 정답 내용을 알고 있는 파일이다.
"""
from __future__ import annotations

import hashlib
import logging
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

__all__ = ["ATTACHMENT_TYPES", "AttachmentError", "font_path", "render_attachment"]

ATTACHMENT_TYPES = ("txt", "docx", "pdf", "image")
PAGE_BREAK = "[다음 쪽]"
_FONT_CANDIDATES = ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/NotoSansKR-VF.ttf",
                    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                    "/System/Library/Fonts/AppleSDGothicNeo.ttc")
_SAFE_NAME = re.compile(r"[^0-9A-Za-z가-힣._ -]")
_EXT = {"txt": ".txt", "docx": ".docx", "pdf": ".pdf", "image": ".png"}
logging.getLogger("fontTools").setLevel(logging.ERROR)   # PDF 글꼴 일부만 넣을 때의 안내 문구를 숨긴다


class AttachmentError(ValueError):
    pass


def font_path() -> str:
    """한글 글꼴 — ``QUESTION_FONT`` 환경 변수, 없으면 흔한 위치에서 찾는다."""
    for candidate in (os.environ.get("QUESTION_FONT"),) + _FONT_CANDIDATES:
        if candidate and Path(candidate).is_file():
            return candidate
    raise AttachmentError("한글 글꼴을 찾지 못했습니다. QUESTION_FONT 환경 변수에 .ttf 경로를 지정해 주세요.")


def _lines(content: str) -> List[str]:
    return [line.rstrip() for line in (content or "").replace("\r\n", "\n").split("\n")]


def _cells(line: str) -> Optional[List[str]]:
    text = line.strip()
    if not (text.startswith("|") and text.endswith("|")):
        return None
    cells = [cell.strip() for cell in text[1:-1].split("|")]
    return None if all(re.fullmatch(r":?-{2,}:?", c or "--") for c in cells) and any(cells) and all(set(c) <= set("-: ") for c in cells) else cells


def _blocks(content: str) -> List[Dict[str, Any]]:
    """줄 표기 → [{kind: heading|bullet|table|page|text, ...}] — 표는 이어진 '|' 줄을 한 덩어리로 묶는다."""
    blocks: List[Dict[str, Any]] = []
    for line in _lines(content):
        stripped = line.strip()
        if stripped == PAGE_BREAK:
            blocks.append({"kind": "page"})
            continue
        row = _cells(line)
        if row is not None:
            if blocks and blocks[-1]["kind"] == "table":
                blocks[-1]["rows"].append(row)
            else:
                blocks.append({"kind": "table", "rows": [row]})
            continue
        if stripped.startswith(("|",)) and set(stripped) <= set("|-: "):
            continue   # Markdown 표 구분선
        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            blocks.append({"kind": "heading", "level": len(heading.group(1)), "text": heading.group(2)})
        elif stripped.startswith(("- ", "* ", "• ")):
            blocks.append({"kind": "bullet", "text": stripped[2:].strip()})
        else:
            blocks.append({"kind": "text", "text": line})
    return blocks


def _write_txt(content: str, path: Path) -> None:
    path.write_text(content, encoding="utf-8")


def _write_docx(content: str, path: Path) -> None:
    import docx   # python-docx

    document = docx.Document()
    for block in _blocks(content):
        if block["kind"] == "page":
            document.add_page_break()
        elif block["kind"] == "heading":
            document.add_heading(block["text"], level=block["level"])
        elif block["kind"] == "bullet":
            document.add_paragraph(block["text"], style="List Bullet")
        elif block["kind"] == "table":
            rows = block["rows"]
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(width):
                    table.cell(r, c).text = row[c] if c < len(row) else ""
        else:
            document.add_paragraph(block["text"])
    document.save(str(path))


def _slice_original(text: str, lines: List[str], start: int, end: int) -> str:
    """줄로 나뉜 글의 [start, end) 줄에 해당하는 원문 구간. 줄 글자를 원문에서 차례로 찾아 자르므로 원문 띄어쓰기가 보존된다."""
    cursor, spans = 0, []
    for piece in lines:
        found = text.find(piece.strip(), cursor) if piece.strip() else -1
        if found < 0:
            return " ".join(p.strip() for p in lines[start:end])   # 못 찾으면 줄을 이어 쓴다 (예외 상황)
        spans.append((found, found + len(piece.strip())))
        cursor = found + len(piece.strip())
    chosen = spans[start:end]
    return text[chosen[0][0]:chosen[-1][1]].strip() if chosen else ""


def _write_pdf(content: str, path: Path) -> str:
    """PDF 를 만들고 평가 모델이 받을 내용을 돌려준다 — 표 칸은 자르지 않고 줄을 바꾸며, 쪽 표시('[n쪽]')는 자동 쪽 넘김까지
    포함해 실제로 놓인 쪽을 따른다 (Codex 리뷰: 칸 40자 자름·자동 쪽 넘김 미반영)."""
    from fpdf import FPDF   # fpdf2

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("KR", fname=font_path())
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    out: List[str] = []
    shown = [0]

    def mark() -> None:   # 지금 쓰는 쪽이 바뀌었으면 쪽 표시를 넣는다
        if pdf.page_no() != shown[0]:
            shown[0] = pdf.page_no()
            out.append(f"[{shown[0]}쪽]")

    def place(height: float) -> None:   # 이 높이가 남은 자리에 안 들어가면 다음 쪽부터 — 쪽 표시가 실제 출력 위치를 따른다
        if pdf.will_page_break(height):
            pdf.add_page()
        mark()

    page_height = pdf.h - pdf.t_margin - pdf.b_margin

    def write_text(text: str, line: float, shown_text: str) -> None:
        """글 한 덩이를 쓴다. 한 쪽에 들어가면 통째로 한 쪽에 두고, 쪽보다 길면 줄마다 쪽을 확인해 쓴다 (Codex 리뷰: 자동 쪽 넘김으로
        갈라진 문단이 평가용 원문에서는 앞 쪽에 적히던 문제)."""
        lines = pdf.multi_cell(width, line, text or " ", dry_run=True, output="LINES")
        if len(lines) * line <= page_height:
            place(len(lines) * line)
            pdf.multi_cell(width, line, text or " ", new_x="LMARGIN", new_y="NEXT")
            out.append(shown_text)
            return
        for n, piece in enumerate(lines):
            place(line)
            pdf.multi_cell(width, line, piece or " ", new_x="LMARGIN", new_y="NEXT")
            out.append((shown_text[:len(shown_text) - len(text)] + piece) if n == 0 else piece)

    for block in _blocks(content):
        if block["kind"] == "page":
            pdf.add_page()
            continue
        if block["kind"] == "heading":
            pdf.set_font("KR", size={1: 16, 2: 14, 3: 12}[block["level"]])
            write_text(block["text"], 9, "#" * block["level"] + " " + block["text"])
            pdf.ln(1)
        elif block["kind"] == "table":
            pdf.set_font("KR", size=9)
            rows = block["rows"]
            cols = max(len(row) for row in rows)
            col_width, line = width / cols, 6
            per_page = max(1, int(page_height // line))
            for row in rows:
                cells = [row[c] if c < len(row) else "" for c in range(cols)]
                cell_lines = [pdf.multi_cell(col_width - 2, line, cell or " ", dry_run=True, output="LINES") or [" "] for cell in cells]
                total = max(len(lines) for lines in cell_lines)
                # 한 쪽보다 긴 행은 줄 단위로 쪽마다 나눠 쓰고, 평가용 원문도 같은 자리에서 나눈다 (Codex 리뷰: 시작 쪽만 적히던 문제).
                # 나뉜 조각의 글은 원문 칸 글자를 줄 위치로 잘라 쓴다 — 줄을 공백으로 이으면 낱말 가운데 공백이 생긴다.
                for offset in range(0, total, per_page):
                    pieces = ["\n".join(lines[offset:offset + per_page]) for lines in cell_lines]
                    height = max(1, max(len(lines[offset:offset + per_page]) for lines in cell_lines)) * line
                    place(height)
                    x, y = pdf.l_margin, pdf.get_y()
                    for i, piece in enumerate(pieces):
                        pdf.rect(x + i * col_width, y, col_width, height)
                        pdf.set_xy(x + i * col_width + 1, y)
                        pdf.multi_cell(col_width - 2, line, piece or " ", new_x="RIGHT", new_y="TOP")
                    pdf.set_xy(x, y + height)
                    shown_cells = cells if total <= per_page else [_slice_original(cell, lines, offset, offset + per_page)
                                                                   for cell, lines in zip(cells, cell_lines)]
                    out.append("| " + " | ".join(shown_cells) + " |")
            pdf.ln(2)
        else:
            pdf.set_font("KR", size=10)
            text = ("• " + block["text"]) if block["kind"] == "bullet" else block["text"]
            write_text(text, 6, ("- " + block["text"]) if block["kind"] == "bullet" else block["text"])
    pdf.output(str(path))
    return "\n".join(out).strip()


def _write_image(spec: Dict[str, Any], path: Path) -> str:
    """PNG — 막대그래프(chart) 또는 글·표. 돌려주는 값은 그림에 실제로 그린 내용(글)."""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(font_path(), 22)
    small = ImageFont.truetype(font_path(), 18)
    chart = spec.get("chart") if isinstance(spec.get("chart"), dict) else None
    if chart:
        labels = [str(x) for x in chart.get("labels") or []]
        values = [float(v) for v in chart.get("values") or []]
        if not labels or len(labels) != len(values):
            raise AttachmentError(f"{spec.get('name')}: chart 의 labels·values 개수가 맞지 않습니다.")
        unit, title = str(chart.get("unit") or ""), str(chart.get("title") or "")
        image = Image.new("RGB", (900, 560), "white")
        draw = ImageDraw.Draw(image)
        draw.text((40, 20), title, fill="black", font=font)
        top, bottom, left = 90, 470, 80
        peak = max(values) or 1.0
        bar = (900 - left - 60) / len(values)
        for i, (label, value) in enumerate(zip(labels, values)):
            x0 = left + i * bar + bar * 0.2
            height = (bottom - top) * value / peak
            draw.rectangle([x0, bottom - height, x0 + bar * 0.6, bottom], fill=(70, 120, 200))
            shown = f"{value:g}{unit}"
            draw.text((x0, bottom - height - 28), shown, fill="black", font=small)
            draw.text((x0, bottom + 10), label, fill="black", font=small)
        draw.line([left, bottom, 880, bottom], fill="black", width=2)
        image.save(str(path))
        return (f"{title}\n" if title else "") + "\n".join(f"{label}: {value:g}{unit}" for label, value in zip(labels, values))
    lines = _lines(spec.get("content") or "")
    height = max(200, 60 + 34 * len(lines))
    image = Image.new("RGB", (1000, height), "white")
    draw = ImageDraw.Draw(image)
    y = 30
    for line in lines:
        row = _cells(line)
        if row is not None:
            col = 920 / max(1, len(row))
            for c, cell in enumerate(row):
                draw.rectangle([40 + c * col, y - 4, 40 + (c + 1) * col, y + 30], outline="black")
                draw.text((48 + c * col, y), cell, fill="black", font=small)
        else:
            draw.text((40, y), line, fill="black", font=font)
        y += 34
    image.save(str(path))
    return "\n".join(lines)


def render_attachment(spec: Dict[str, Any], out_dir: Path) -> Dict[str, Any]:
    """명세 → 파일을 만들고 attachments 항목 {name, type, content, sha256, size, path} 를 돌려준다 (path 는 out_dir 기준 파일 이름)."""
    if not isinstance(spec, dict):
        raise AttachmentError("첨부 명세는 객체여야 합니다.")
    kind = str(spec.get("type") or "").strip().lower()
    if kind not in ATTACHMENT_TYPES:
        raise AttachmentError(f"첨부 종류는 {', '.join(ATTACHMENT_TYPES)} 중 하나여야 합니다 ({kind or '없음'}).")
    name = _SAFE_NAME.sub("_", str(spec.get("name") or "").strip())[:80]
    if not name:
        raise AttachmentError("첨부 이름이 비었습니다.")
    if not name.lower().endswith(_EXT[kind]):
        name = Path(name).stem + _EXT[kind]
    content = spec.get("content") if isinstance(spec.get("content"), str) else ""
    if kind != "image" and not content.strip():
        raise AttachmentError(f"{name}: 첨부 내용(content)이 비었습니다.")
    out_dir.mkdir(parents=True, exist_ok=True)
    temp = out_dir / f".tmp-{os.getpid()}-{hashlib.sha1(name.encode('utf-8')).hexdigest()[:8]}{_EXT[kind]}"
    pdf_text = ""
    if kind == "txt":
        _write_txt(content, temp)
    elif kind == "docx":
        _write_docx(content, temp)
    elif kind == "pdf":
        pdf_text = _write_pdf(content, temp)
    else:
        content = _write_image(spec, temp)
    data = temp.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    final = out_dir / f"{digest}{_EXT[kind]}"
    if final.exists():
        temp.unlink()
    else:
        temp.replace(final)
    text = pdf_text if kind == "pdf" else judge_text(content, kind)   # PDF 는 실제로 그린 쪽 배치를 따른다
    return {"name": name, "type": kind, "content": text, "sha256": digest, "size": len(data), "path": final.name}


def judge_text(content: str, kind: str) -> str:
    """평가 모델이 받는 첨부 내용 — 표 구분선은 빼고, PDF·DOCX 는 쪽 표시('[1쪽]'·'[2쪽]' …)를 남긴다 (J8-2 쪽 읽기)."""
    lines = [line for line in _lines(content) if not (line.strip().startswith("|") and set(line.strip()) <= set("|-: "))]
    if kind not in ("pdf", "docx"):
        return "\n".join(lines).strip()
    out, page = ["[1쪽]"], 1
    for line in lines:
        if line.strip() == PAGE_BREAK:
            page += 1
            out.append(f"[{page}쪽]")
        else:
            out.append(line)
    return "\n".join(out).strip()
